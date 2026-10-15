import asyncio
from datetime import datetime, timedelta, timezone
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

import app.service as service_module
import app.stock_ops as stock_module
from app.config import CabinetConfig, Settings
from app.distributor import distribute_barcode
from app.excel_io import read_input_with_controls
from app.models import ProductVariant, Warehouse
from app.service import DistributionService
from app.stock_ops import StockManager


def workbook(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    book.active.append(header)
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


class Session:
    def __init__(self, **kwargs):
        pass

    async def __aenter__(self):
        return self

    async def __aexit__(self, *args):
        return False


class FakeWB:
    def __init__(self, cabinet, session):
        self.cabinet = cabinet

    async def get_warehouses(self):
        return [Warehouse(self.cabinet.key, self.cabinet.name, ident, f'Warehouse {ident}')
                for ident in (11, 12)]

    async def get_catalog_variants(self):
        offset = 0 if self.cabinet.key == 'cabinet_1' else 100
        variants = {
            offset + 1: ProductVariant(self.cabinet.key, offset + 1, 1, ('single', 'single_alias')),
            offset + 2: ProductVariant(self.cabinet.key, offset + 2, 2, ('kit', 'kit_alias')),
            offset + 3: ProductVariant(self.cabinet.key, offset + 3, 3, ('other_kit',)),
        }
        if self.cabinet.key == 'cabinet_2':
            variants[offset + 2] = ProductVariant(self.cabinet.key, offset + 2, 2, ('kit_alias',))
        return {sku: variant for variant in variants.values() for sku in variant.skus}, variants

    async def get_orders(self, days):
        assert days == 14
        offset = 0 if self.cabinet.key == 'cabinet_1' else 100
        return [{'chrtId': offset + ident, 'warehouseId': warehouse}
                for ident in (1, 2, 3) for warehouse in (11, 12) for order in range(10)]

    async def get_fbs_stocks(self, warehouse, ids):
        return {int(ident): self.stock.get((self.cabinet.key, warehouse, int(ident)), 7) for ident in ids}

    async def put_fbs_stocks(self, warehouse, amounts):
        self.puts.append((self.cabinet.key, warehouse, dict(amounts)))
        for ident, amount in amounts.items():
            self.stock[self.cabinet.key, warehouse, ident] = amount


@pytest.fixture
def project(tmp_path, monkeypatch):
    settings = Settings('fake', {123}, (
        CabinetConfig('cabinet_1', 'TT', 'fake1'),
        CabinetConfig('cabinet_2', 'Group', 'fake2'),
    ), 14, 4, 2, 20, 4, base_dir=tmp_path, admin_ids={123})
    workbook(tmp_path / 'templates' / 'stocks_template.xlsx', ['Баркод', 'Количество'], [])
    workbook(tmp_path / 'data' / 'kits_template.xlsx',
             ['Название', 'баркод комплекта', 'баркод1', 'баркод2'],
             [('Kit', 'kit', 'single', 'single'), ('Other', 'other_kit', 'single', 'single')])
    monkeypatch.setattr(service_module.aiohttp, 'TCPConnector', lambda **kwargs: None)
    monkeypatch.setattr(service_module.aiohttp, 'ClientSession', Session)
    monkeypatch.setattr(service_module, 'WBClient', FakeWB)
    monkeypatch.setattr(stock_module, 'WBClient', FakeWB)
    FakeWB.stock = {}
    FakeWB.puts = []
    return DistributionService(settings)


def process(project, rows, header='Доступно'):
    source = workbook(project.settings.base_dir / 'input.xlsx', ['Код', header], rows)
    return asyncio.run(project.process(source, 123))


def allocation(result, *barcodes):
    return sum(sum(line['allocations'].values()) for line in result['report']['allocations']
               if line['barcode'] in barcodes)


def warehouse_rows(result):
    output = []
    for path in result['files']:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            output.append(dict(book.active.iter_rows(min_row=2, values_only=True)))
        finally:
            book.close()
    return output


def test_positive_ready_kit_does_not_consume_components(project):
    result = process(project, [('single', 80), ('kit', 16), ('other_kit', '!')])
    assert allocation(result, 'kit') == 16
    assert allocation(result, 'single') == 80
    assert result['report']['component_consumed_by_kits'] == {}
    assert result['summary'].kit_units == 0
    assert all(rows.get('kit', rows.get('kit_alias')) == 4 for rows in warehouse_rows(result))


@pytest.mark.parametrize('kit_row', [[], [('kit', 0)]])
def test_zero_or_absent_kit_is_still_assembled(project, kit_row):
    result = process(project, [('single', 80), ('other_kit', '!')] + kit_row)
    built = allocation(result, 'kit')
    assert built >= 16
    assert result['report']['component_consumed_by_kits']['single'] == built * 2
    assert allocation(result, 'single') + built * 2 + result['summary'].reserve_units == 80


def test_percent_writes_zero_in_all_warehouse_files(project):
    date_before = datetime.now(timezone(timedelta(hours=3))).strftime('%d.%m.%Y')
    result = process(project, [('single', 80), ('kit', '%'), ('other_kit', '!')])
    date_after = datetime.now(timezone(timedelta(hours=3))).strftime('%d.%m.%Y')
    assert result['zip'].name in {f'Распределение_FBS_{date}.zip' for date in (date_before, date_after)}
    assert allocation(result, 'kit', 'kit_alias') == 0
    assert allocation(result, 'single') == 80
    assert result['report']['component_consumed_by_kits'] == {}
    for rows in warehouse_rows(result):
        assert rows.get('kit', rows.get('kit_alias')) == 0
        assert 'other_kit' not in rows
    with ZipFile(result['zip']) as archive:
        assert set(archive.namelist()) == {path.name for path in result['files']} | {'report.json'}


def test_positive_alternative_kit_barcode_prevents_assembly(project):
    result = process(project, [('single', 80), ('kit_alias', 9), ('kit', 7), ('other_kit', '!')])
    assert allocation(result, 'kit', 'kit_alias') == 16
    assert allocation(result, 'single') == 80
    assert not result['report']['component_consumed_by_kits']


@pytest.mark.parametrize('rows', [
    [('kit', '%'), ('kit_alias', 16)],
    [('kit_alias', 16), ('kit', '%')],
    [('kit', 16), ('kit', '%')],
    [('kit', '%'), ('kit', 16)],
])
def test_percent_overrides_positive_stock_independent_of_order(project, rows):
    result = process(project, [('single', 80), ('other_kit', '!')] + rows)
    assert allocation(result, 'kit', 'kit_alias') == 0
    assert allocation(result, 'single') == 80
    for exported in warehouse_rows(result):
        assert exported.get('kit', exported.get('kit_alias')) == 0


@pytest.mark.parametrize('quantity', [16, '%'])
def test_disabled_assembly_releases_components_for_other_kits(project, quantity):
    result = process(project, [('single', 100), ('kit', quantity)])
    assembled = {item['barcode']: item['built'] for item in result['report']['kits']}
    assert assembled.get('kit', 0) == 0
    assert assembled['other_kit'] > 0
    assert allocation(result, 'single') > 0
    assert (allocation(result, 'single') + assembled['other_kit'] * 2
            + result['summary'].reserve_units == 100)


@pytest.mark.parametrize('rows', [
    [('kit', '!'), ('kit', '%')],
    [('kit', '%'), ('kit', '!')],
    [('kit', '%'), ('kit_alias', '!')],
])
def test_manual_exclusion_takes_priority_over_percent(project, rows):
    result = process(project, [('single', 80), ('other_kit', '!')] + rows)
    for exported in warehouse_rows(result):
        assert 'kit' not in exported and 'kit_alias' not in exported


def test_cabinet_exclusion_keeps_priority(project):
    workbook(project.exclusions_path, ['ШК', 'ID кабинета 1', 'ID кабинета 2'],
             [('kit_alias', None, 'seller2')])
    result = process(project, [('single', 80), ('kit', '%'), ('other_kit', '!')])
    for path, rows in zip(result['files'], warehouse_rows(result)):
        if path.name.startswith('Group'):
            assert 'kit_alias' not in rows
        else:
            assert rows['kit'] == 0


def test_percent_reaches_automatic_upload(project):
    result = process(project, [('single', 80), ('kit', '%'), ('other_kit', '!')])
    manager = StockManager(project.settings)
    async def scenario():
        preview = await manager.preview('distribution', 123, {'cabinet_1', 'cabinet_2'}, files=result['files'])
        assert not FakeWB.puts
        for key, rows in preview['targets'].items():
            ident = '2' if key.startswith('cabinet_1:') else '102'
            assert rows[ident] == 0
        record = await manager.execute(preview['id'], 123)
        assert record['status'] == 'verified'
        for cabinet in ('cabinet_1', 'cabinet_2'):
            ident = 2 if cabinet == 'cabinet_1' else 102
            for warehouse in (11, 12):
                assert FakeWB.stock[cabinet, warehouse, ident] == 0
    asyncio.run(scenario())


def test_availability_header_and_spaced_marker(project):
    result = process(project, [('single', 80), ('kit', ' % '), ('other_kit', '!')], header='Наличие')
    assert allocation(result, 'kit', 'kit_alias') == 0
    assert allocation(result, 'single') == 80


@pytest.mark.parametrize('quantity', [1, 3, 17, 23])
def test_ready_kit_uses_existing_single_distribution_rules(project, quantity):
    result = process(project, [('single', 80), ('kit', quantity), ('other_kit', '!')])
    expected = distribute_barcode('kit', quantity,
        [Warehouse(cabinet.key, cabinet.name, ident, f'Warehouse {ident}')
         for cabinet in project.settings.cabinets for ident in (11, 12)],
        {cabinet.key: {'kit': ProductVariant(cabinet.key, ident, 2, ('kit',))}
         for cabinet, ident in zip(project.settings.cabinets, (2, 102))},
        {'cabinet_1': {(2, 11): 10, (2, 12): 10},
         'cabinet_2': {(102, 11): 10, (102, 12): 10}}, 20, 2, 4)
    actual = next(line for line in result['report']['allocations'] if line['barcode'] == 'kit')
    assert actual['allocations'] == {f'{cabinet}:{warehouse}': amount
                                     for (cabinet, warehouse), amount in expected.allocations.items() if amount}
    assert actual['reserve_qty'] == expected.reserve_qty
    assert not result['report']['component_consumed_by_kits']


def test_ready_kit_without_orders_keeps_extra_in_reserve(project, monkeypatch):
    async def no_orders(client, days):
        return []
    monkeypatch.setattr(FakeWB, 'get_orders', no_orders)
    result = process(project, [('single', 80), ('kit', 30), ('other_kit', '!')])
    assert allocation(result, 'kit') == 16
    actual = next(line for line in result['report']['allocations'] if line['barcode'] == 'kit')
    assert actual['reserve_qty'] == 14
    assert not result['report']['component_consumed_by_kits']


def test_percent_for_unknown_kit_is_rejected(project):
    with pytest.raises(ValueError, match='не найденных в шаблоне комплектов'):
        process(project, [('single', '%')])
    assert not FakeWB.puts


def test_percent_only_file_is_valid(project):
    result = process(project, [('kit', '%')])
    assert result['summary'].input_units == 0
    assert result['summary'].kit_units == 0
    assert all(not any(rows.values()) for rows in warehouse_rows(result))


def test_controls_do_not_persist_to_next_calculation(project):
    first = process(project, [('single', 80), ('kit', '%'), ('other_kit', '!')])
    assert allocation(first, 'kit') == 0
    second = process(project, [('single', 80), ('kit', 0), ('other_kit', '!')])
    assert allocation(second, 'kit') > 0


def test_xls_marker_uses_same_reader_rules(tmp_path, monkeypatch):
    import xlrd
    rows = [['Код', 'Наличие'], ['kit', ' % '], ['single', 80], ['other', '!']]
    class Sheet:
        nrows = len(rows)
        ncols = 2

        def cell_value(self, row, column):
            return rows[row][column]

    class Book:
        def sheet_by_index(self, index):
            return Sheet()

    monkeypatch.setattr(xlrd, 'open_workbook', lambda path: Book())
    stock, excluded, zero_kits = read_input_with_controls(tmp_path / 'input.xls')
    assert stock == {'kit': 0, 'single': 80}
    assert excluded == {'other'}
    assert zero_kits == {'kit'}


def test_percent_conflicting_catalog_ids_are_rejected(project, monkeypatch):
    original_catalog = FakeWB.get_catalog_variants
    async def conflicting_catalog(client):
        by_sku, by_chrt = await original_catalog(client)
        if client.cabinet.key == 'cabinet_1':
            variant = ProductVariant(client.cabinet.key, 4, 4, ('kit_alias',))
            by_chrt[2] = ProductVariant(client.cabinet.key, 2, 2, ('kit',))
            by_sku['kit'] = by_chrt[2]
            by_chrt[4] = variant
            by_sku['kit_alias'] = variant
        else:
            variant = ProductVariant(client.cabinet.key, 102, 2, ('kit', 'kit_alias'))
            by_chrt[102] = variant
            by_sku.update({'kit': variant, 'kit_alias': variant})
        return by_sku, by_chrt
    monkeypatch.setattr(FakeWB, 'get_catalog_variants', conflicting_catalog)
    with pytest.raises(ValueError, match='противоречивые ШК'):
        process(project, [('single', 80), ('kit', '%'), ('kit_alias', 16)])
    assert not FakeWB.puts
