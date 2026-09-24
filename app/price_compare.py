"""Independent read-only comparison of seller prices and WB orders for two cabinets.

This module does not change the distribution service, inventory, or price settings.
"""
from __future__ import annotations

import asyncio
import os
from dataclasses import replace
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import aiohttp
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Settings
from .wb_api import WBClient, WBApiError

PRICE_BASE = "https://discounts-prices-api.wildberries.ru"
STATS_BASE = "https://statistics-api.wildberries.ru"
MSK = ZoneInfo("Europe/Moscow")


def _seller_price_for_variant(good: dict, chrt_id: int) -> float | None:
    """Use discountedPrice (seller discount included, WB discounts excluded)."""
    sizes = good.get("sizes") or []
    if not sizes:
        return None
    matched = [s for s in sizes if str(s.get("sizeID")) == str(chrt_id)]
    if not matched:
        # No ambiguous price attribution if a card has several differently-priced sizes.
        values = {s.get("discountedPrice") for s in sizes if s.get("discountedPrice") is not None}
        if len(values) != 1:
            return None
        matched = sizes[:1]
    value = matched[0].get("discountedPrice")
    return float(value) if value is not None else None


async def _load_prices(client: WBClient) -> dict[int, dict]:
    by_nm: dict[int, dict] = {}
    offset = 0
    while True:
        response = await client._request(
            "GET", f"{PRICE_BASE}/api/v2/list/goods/filter",
            params={"limit": 1000, "offset": offset},
        )
        goods = ((response or {}).get("data") or {}).get("listGoods") or []
        for good in goods:
            if good.get("nmID") is not None:
                by_nm[int(good["nmID"])] = good
        if len(goods) < 1000:
            break
        offset += 1000
    return by_nm


async def _load_statistics_orders(client: WBClient, days: int) -> list[dict]:
    # flag=0 includes subsequent updates; filter by original order date below.
    # Last N *complete* Moscow calendar days; exclude today, so a weekly and
    # fortnightly report compare like-for-like periods across both cabinets.
    today = datetime.now(MSK).date()
    start = today - timedelta(days=days)
    response = await client._request(
        "GET", f"{STATS_BASE}/api/v1/supplier/orders",
        params={"dateFrom": start.isoformat(), "flag": 0},
    )
    if not isinstance(response, list):
        raise WBApiError(f"{client.cabinet.name}: неожиданный формат статистики заказов")
    if len(response) >= 79000:
        raise WBApiError(
            f"{client.cabinet.name}: статистика содержит {len(response)} строк, "
            "возможна неполная выгрузка (лимит API ~80 000). Отчёт не сформирован."
        )
    # Each srid is one ordered item; latest update determines cancellation.
    latest: dict[str, dict] = {}
    for idx, order in enumerate(response):
        order_date = str(order.get("date") or "")[:10]
        if not (start.isoformat() <= order_date < today.isoformat()):
            continue
        uid = str(order.get("srid") or f"row:{idx}")
        previous = latest.get(uid)
        if previous is None or str(order.get("lastChangeDate") or "") >= str(previous.get("lastChangeDate") or ""):
            latest[uid] = order
    return [order for order in latest.values() if not order.get("isCancel")]


def _match_variants(catalogs: list[dict[int, Any]]) -> tuple[list[tuple[Any, Any]], int]:
    """Match only SKUs shared between both cabinets; ignore ambiguous mappings."""
    first, second = catalogs
    sku_to_second: dict[str, set[int]] = defaultdict(set)
    for variant in second.values():
        for sku in variant.skus:
            sku_to_second[str(sku)].add(variant.chrt_id)
    candidates: dict[int, set[int]] = {}
    for variant in first.values():
        matched = set().union(*(sku_to_second.get(str(sku), set()) for sku in variant.skus))
        if matched:
            candidates[variant.chrt_id] = matched
    reverse: dict[int, set[int]] = defaultdict(set)
    for left, rights in candidates.items():
        for right in rights:
            reverse[right].add(left)
    pairs = [(first[left], second[next(iter(rights))]) for left, rights in candidates.items()
             if len(rights) == 1 and len(reverse[next(iter(rights))]) == 1]
    ambiguous = sum(1 for rights in candidates.values() if len(rights) != 1 or len(reverse[next(iter(rights))]) != 1)
    return pairs, ambiguous


def _order_counts(orders: list[dict], by_sku: dict[str, Any], by_chrt: dict[int, Any]) -> dict[int, int]:
    counts: dict[int, int] = defaultdict(int)
    for order in orders:
        variant = by_sku.get(str(order.get("barcode") or "").strip())
        if variant is None:
            nm = order.get("nmId")
            # Historical barcodes may be missing in current catalog; only match
            # on nmID when that nmID uniquely identifies one current variation.
            if nm is not None:
                matches = [v for v in by_chrt.values() if str(v.nm_id) == str(nm)]
                if len(matches) == 1:
                    variant = matches[0]
        if variant is not None:
            counts[variant.chrt_id] += 1
    return counts


def _export_report(path: Path, rows: list[tuple], names: list[str], days: int, ambiguous: int,
                   start_date: str, end_date: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение цен"
    ws.append([f"Сравнение цен и заказов • {start_date} — {end_date} • {days} дней"])
    ws.merge_cells("A1:J1")
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="193B63")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.append(["Текущие цены продавца на момент выгрузки; заказы FBO + FBS, без отмен; только товары, присутствующие в обоих кабинетах."])
    ws.merge_cells("A2:J2")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 34
    ws.append([
        "ШК (общий)", "Артикул продавца", f"Цена {names[0]}, ₽", f"Заказы {names[0]}, шт.",
        f"Цена {names[1]}, ₽", f"Заказы {names[1]}, шт.", "Разница заказов, %",
        "Разница цен, %", "ШК кабинета 1", "ШК кабинета 2",
    ])
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315C83")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 40
    for row in rows:
        ws.append(row)
        idx = ws.max_row
        # Formula-driven differences; zero baseline remains blank, not infinity.
        ws.cell(idx, 7, f'=IF(OR(D{idx}="",D{idx}=0),"",(F{idx}-D{idx})/D{idx})')
        ws.cell(idx, 8, f'=IF(OR(C{idx}="",C{idx}=0,E{idx}=""),"",(E{idx}-C{idx})/C{idx})')
        for col in (1, 9, 10):
            ws.cell(idx, col).number_format = "@"
        for col in (3, 5):
            ws.cell(idx, col).number_format = '#,##0.00'
        for col in (7, 8):
            ws.cell(idx, col).number_format = '+0.0%;-0.0%;0.0%'
        if idx % 2 == 0:
            for cell in ws[idx]:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
    ws.auto_filter.ref = f"A3:J{max(3, ws.max_row)}"
    ws.freeze_panes = "C4"
    widths = [22, 27, 19, 21, 19, 21, 21, 19, 22, 22]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    info = wb.create_sheet("Примечания")
    for record in [
        ("Период", f"{start_date} — {end_date}; {days} полных календарных дней (МСК)"),
        ("Цена", "Текущая discountedPrice: цена продавца с его скидкой, без скидок WB/клуба."),
        ("Заказы", "Заказы FBO и FBS из статистики WB; отменённые исключены. Возможна задержка данных API."),
        ("Разница заказов", "(Заказы 2 − Заказы 1) / Заказы 1; при нуле в кабинете 1 — пусто."),
        ("Разница цен", "(Цена 2 − Цена 1) / Цена 1; при нуле/нет цены — пусто."),
        ("Сопоставление", "Только общие ШК двух кабинетов; альтернативные ШК одного размера объединяются."),
        ("Артикул продавца", "Из первого кабинета; при отсутствии — из второго."),
        ("Неоднозначные совпадения", f"Не включены товарные вариации с неоднозначным сопоставлением: {ambiguous}."),
        ("Ограничение", "Текущую цену нельзя считать исторической ценой за все 7/14 дней; корреляция не доказывает причинность."),
    ]:
        info.append(record)
    info.column_dimensions["A"].width = 29
    info.column_dimensions["B"].width = 115
    for row in info:
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


async def build_price_comparison(settings: Settings, days: int, output: Path) -> dict:
    if days not in (7, 14):
        raise ValueError("Период сравнения может быть только 7 или 14 дней")
    async with aiohttp.ClientSession() as session:
        clients = [WBClient(cabinet, session) for cabinet in settings.cabinets]
        async def get_data(client: WBClient):
            by_sku, by_chrt = await client.get_catalog_variants()
            index = 1 if client.cabinet.key == settings.cabinets[0].key else 2
            price_token = os.getenv(f"WB_PRICES_TOKEN_{index}", "").strip()
            stats_token = os.getenv(f"WB_STATS_TOKEN_{index}", "").strip()
            price_client = (WBClient(replace(client.cabinet, token=price_token), session)
                            if price_token else client)
            stats_client = (WBClient(replace(client.cabinet, token=stats_token), session)
                            if stats_token else client)
            prices = await _load_prices(price_client)
            orders = await _load_statistics_orders(stats_client, days)
            return by_sku, by_chrt, prices, _order_counts(orders, by_sku, by_chrt)
        cabinet_data = await asyncio.gather(*(get_data(client) for client in clients))
    (skus_1, catalog_1, price_1, counts_1), (skus_2, catalog_2, price_2, counts_2) = cabinet_data
    pairs, ambiguous = _match_variants([catalog_1, catalog_2])
    report_rows = []
    missing_prices = 0
    for left, right in pairs:
        common = sorted(set(left.skus) & set(right.skus))
        if not common:
            continue
        good_1, good_2 = price_1.get(left.nm_id, {}), price_2.get(right.nm_id, {})
        p1, p2 = _seller_price_for_variant(good_1, left.chrt_id), _seller_price_for_variant(good_2, right.chrt_id)
        if p1 is None or p2 is None:
            missing_prices += 1
        barcode = common[0]
        article = str(good_1.get("vendorCode") or good_2.get("vendorCode") or "")
        report_rows.append((barcode, article, p1, counts_1.get(left.chrt_id, 0),
                            p2, counts_2.get(right.chrt_id, 0), None, None,
                            ", ".join(left.skus), ", ".join(right.skus)))
    report_rows.sort(key=lambda row: (str(row[1]).lower(), row[0]))
    today = datetime.now(MSK).date()
    _export_report(output, report_rows, [c.name for c in settings.cabinets], days, ambiguous,
                   (today - timedelta(days=days)).isoformat(), (today - timedelta(days=1)).isoformat())
    return {"products": len(report_rows), "missing_prices": missing_prices, "ambiguous": ambiguous,
            "orders_1": sum(row[3] for row in report_rows), "orders_2": sum(row[5] for row in report_rows)}
