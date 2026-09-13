from __future__ import annotations

import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

from .models import DistributionLine, RunSummary, Warehouse


BARCODE_HEADERS = {"код", "баркод", "штрихкод", "barcode", "шк"}
QTY_HEADERS = {"доступно", "количество", "кол-во", "qty", "quantity", "остаток"}


def _norm(value) -> str:
    return str(value or "").strip().lower().replace("ё", "е")


def _iter_input_rows(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        headers = [_norm(ws.cell(1, col).value) for col in range(1, ws.max_column + 1)]
        yield headers
        for row in range(2, ws.max_row + 1):
            yield [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        return
    if suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path)
        ws = wb.sheet_by_index(0)
        if ws.nrows == 0:
            return
        yield [_norm(ws.cell_value(0, col)) for col in range(ws.ncols)]
        for row in range(1, ws.nrows):
            yield [ws.cell_value(row, col) for col in range(ws.ncols)]
        return
    raise ValueError("Поддерживаются только файлы .xlsx и .xls")


def read_input_excel(path: Path) -> tuple[dict[str, int], set[str]]:
    rows = iter(_iter_input_rows(path))
    try:
        headers = next(rows)
    except StopIteration:
        raise ValueError("Файл пустой")

    barcode_col = next((i for i, h in enumerate(headers) if h in BARCODE_HEADERS), None)
    qty_col = next((i for i, h in enumerate(headers) if h in QTY_HEADERS), None)
    if barcode_col is None or qty_col is None:
        raise ValueError(
            "В файле должны быть две колонки: «Код» и «Доступно» "
            "(также поддерживаются старые «Баркод» и «Количество»)"
        )

    result: dict[str, int] = {}
    excluded: set[str] = set()
    for values in rows:
        barcode_raw = values[barcode_col] if barcode_col < len(values) else None
        qty_raw = values[qty_col] if qty_col < len(values) else None
        if barcode_raw in (None, "") and qty_raw in (None, ""):
            continue
        barcode = str(barcode_raw or "").strip()
        if barcode.endswith(".0") and barcode[:-2].isdigit():
            barcode = barcode[:-2]
        if not barcode:
            continue
        if str(qty_raw or "").strip() == "!":
            excluded.add(barcode)
            result.pop(barcode, None)
            continue
        if barcode in excluded:
            continue
        try:
            qty = int(float(qty_raw or 0))
        except (TypeError, ValueError):
            raise ValueError(f"Некорректное количество для ШК {barcode}: {qty_raw}")
        if qty < 0:
            raise ValueError(f"Количество не может быть отрицательным: ШК {barcode}")
        result[barcode] = result.get(barcode, 0) + qty
    if not result and not excluded:
        raise ValueError("В файле нет строк с остатками")
    return result, excluded


# Обратная совместимость для внутренних импортов старых версий.
def read_input_xlsx(path: Path) -> tuple[dict[str, int], set[str]]:
    return read_input_excel(path)


def safe_filename(value: str) -> str:
    value = re.sub(r'[<>:"/\\|?*]+', "_", value).strip().strip(".")
    return value[:140] or "warehouse"


def write_warehouse_files(
    template_path: Path,
    output_dir: Path,
    warehouses: list[Warehouse],
    lines: list[DistributionLine],
    all_barcodes_by_cabinet: dict[str, set[str]],
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    files: list[Path] = []

    # Суммируем, а не перезаписываем: баркод комплекта теоретически может
    # одновременно присутствовать и как готовый физический товар во входном файле.
    by_wh: dict[tuple[str, int], dict[str, int]] = {}
    for line in lines:
        for key, qty in line.allocations.items():
            if qty > 0:
                bucket = by_wh.setdefault(key, {})
                bucket[line.barcode] = bucket.get(line.barcode, 0) + int(qty)

    for wh in warehouses:
        allocated = by_wh.get((wh.cabinet_key, wh.warehouse_id), {})

        # WB не обнуляет позиции, отсутствующие в загружаемом файле.
        # Поэтому каждый файл склада содержит ВСЕ доступные баркоды кабинета:
        # распределённым ставим рассчитанный остаток, всем остальным — 0.
        catalog_barcodes = set(all_barcodes_by_cabinet.get(wh.cabinet_key, set()))
        all_rows = catalog_barcodes | set(allocated)

        filename = f"{safe_filename(wh.cabinet_name)} — {safe_filename(wh.name)}.xlsx"
        target = output_dir / filename
        shutil.copy2(template_path, target)
        wb = load_workbook(target)
        ws = wb.active
        ws.title = "Остатки"
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for barcode in sorted(all_rows):
            ws.append([barcode, int(allocated.get(barcode, 0))])
        wb.save(target)
        files.append(target)
    return files


def build_summary(
    lines: list[DistributionLine],
    output_files: int,
    no_sales: list[str],
    not_found: list[str],
    no_sales_kits: list[str],
    not_found_kits: list[str],
    physical_input_units: int,
    excluded_barcodes: set[str] | None = None,
) -> RunSummary:
    summary = RunSummary()
    single_lines = [x for x in lines if not x.is_kit]
    kit_lines = [x for x in lines if x.is_kit]
    summary.input_lines = len(single_lines)
    summary.input_units = physical_input_units
    summary.reserve_units = sum(x.reserve_qty for x in single_lines)
    summary.allocated_units = max(0, summary.input_units - summary.reserve_units)
    summary.output_units = sum(sum(x.allocations.values()) for x in lines)
    summary.kit_units = sum(sum(x.allocations.values()) for x in kit_lines)
    summary.kit_skus = sum(1 for x in kit_lines if sum(x.allocations.values()) > 0)
    summary.output_files = output_files
    summary.no_sales_barcodes = no_sales
    summary.not_found_barcodes = not_found
    summary.no_sales_kit_barcodes = no_sales_kits
    summary.not_found_kit_barcodes = not_found_kits
    summary.excluded_barcodes = sorted(excluded_barcodes or set())
    return summary
