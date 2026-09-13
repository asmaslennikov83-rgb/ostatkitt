from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from .models import KitDefinition


def _norm(value) -> str:
    return str(value or "").strip().lower().replace("ё", "е")


def _barcode(value) -> str:
    value = str(value or "").strip()
    if value.endswith(".0") and value[:-2].isdigit():
        value = value[:-2]
    return value


def _read_table(path: Path) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        return [
            [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
            for row in range(1, ws.max_row + 1)
        ]
    if suffix == ".xls":
        import xlrd
        wb = xlrd.open_workbook(path)
        ws = wb.sheet_by_index(0)
        return [[ws.cell_value(row, col) for col in range(ws.ncols)] for row in range(ws.nrows)]
    raise ValueError("Шаблон комплектов должен быть в формате .xlsx или .xls")


def read_kits_excel(path: Path) -> list[KitDefinition]:
    table = _read_table(path)
    if not table:
        raise ValueError("Шаблон комплектов пустой")

    header_row = table[0]
    headers = {_norm(value): idx for idx, value in enumerate(header_row)}
    kit_col = headers.get("баркод комплекта")
    if kit_col is None:
        raise ValueError("В шаблоне комплектов нет колонки «баркод комплекта»")

    name_col = headers.get("название")
    component_cols: list[tuple[int, int]] = []
    for header, col in headers.items():
        match = re.fullmatch(r"баркод\s*(\d+)", header)
        if match:
            component_cols.append((int(match.group(1)), col))
    component_cols.sort()
    if not component_cols:
        raise ValueError("В шаблоне комплектов должны быть колонки баркод1, баркод2, ...")

    kits: list[KitDefinition] = []
    seen: dict[str, int] = {}
    for row_index, values in enumerate(table[1:], start=2):
        def value_at(idx: int | None):
            if idx is None or idx >= len(values):
                return None
            return values[idx]

        kit_barcode = _barcode(value_at(kit_col))
        components = tuple(
            bc for _idx, col in component_cols
            if (bc := _barcode(value_at(col)))
        )
        name = str(value_at(name_col) or "").strip() if name_col is not None else ""

        if not kit_barcode and not components and not name:
            continue
        if not kit_barcode:
            raise ValueError(f"Строка {row_index}: не указан баркод комплекта")
        if not components:
            raise ValueError(f"Строка {row_index}: для комплекта {kit_barcode} не указаны компоненты")
        if kit_barcode in seen:
            raise ValueError(
                f"Баркод комплекта {kit_barcode} повторяется в строках {seen[kit_barcode]} и {row_index}"
            )
        seen[kit_barcode] = row_index
        kits.append(KitDefinition(name=name, barcode=kit_barcode, components=components, row_number=row_index))

    if not kits:
        raise ValueError("В шаблоне комплектов нет комплектов")
    return kits


def read_kits_xlsx(path: Path) -> list[KitDefinition]:
    return read_kits_excel(path)

