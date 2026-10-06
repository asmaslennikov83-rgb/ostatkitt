"""Audited, opt-in WB FBS stock updates, emergency zeroing and guarded rollback.

No API mutation occurs in preview(). Every mutation requires a fresh preview,
admin confirmation and a fresh compare-and-swap check immediately before PUT.
"""
from __future__ import annotations

import asyncio
import json
import os
import uuid
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import aiohttp
from openpyxl import Workbook, load_workbook

from .config import Settings
from .excel_io import safe_filename
from .exclusions import read_exclusions
from .wb_api import WBClient, WBApiError


class StockOperationError(RuntimeError):
    pass


def utcnow() -> str:
    return datetime.now(timezone.utc).isoformat()


def _save(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix('.tmp')
    with tmp.open('w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    tmp.replace(path)


def _load(path: Path) -> dict:
    return json.loads(path.read_text(encoding='utf-8'))


def _key(cabinet_key: str, warehouse_id: int) -> str:
    return f'{cabinet_key}:{warehouse_id}'


def _chunks(items: dict[int, int], size: int = 500):
    keys = sorted(items)
    for offset in range(0, len(keys), size):
        yield {key: items[key] for key in keys[offset:offset + size]}


def compare(before: dict, targets: dict) -> dict:
    totals = Counter()
    for key, wanted in targets.items():
        old = before[key]
        for chrt, qty in wanted.items():
            prev = old[chrt]
            totals['positions'] += 1
            if prev == qty:
                totals['unchanged'] += 1
            elif qty == 0:
                totals['zeroed'] += 1
            elif qty > prev:
                totals['increased'] += 1
            else:
                totals['decreased'] += 1
            if prev != qty:
                totals['changed'] += 1
    return dict(totals)


def load_distribution_targets(files: list[Path], warehouses: dict, catalogs: dict) -> tuple[dict, dict]:
    """Map exported SKU files to unambiguous cabinet-local chrtIDs."""
    targets, labels = {}, {}
    seen_files = set()
    for key, wh in warehouses.items():
        cab_key = key.split(':', 1)[0]
        name = f'{safe_filename(wh.cabinet_name)} — {safe_filename(wh.name)}.xlsx'
        if name in seen_files:
            raise StockOperationError(f'Совпали имена складских файлов: {name}')
        seen_files.add(name)
        matching = [path for path in files if path.name == name]
        if len(matching) != 1:
            raise StockOperationError(f'Нет однозначного Excel-файла для склада: {name}')
        by_barcode = catalogs[cab_key]
        result = {}
        book = load_workbook(matching[0], read_only=True, data_only=True)
        try:
            for row in book.active.iter_rows(min_row=2, values_only=True):
                barcode = str(row[0] or '').strip()
                if barcode.endswith('.0') and barcode[:-2].isdigit():
                    barcode = barcode[:-2]
                if not barcode:
                    continue
                try:
                    qty = int(row[1])
                    if float(row[1]) != qty or qty < 0:
                        raise ValueError()
                except (ValueError, TypeError):
                    raise StockOperationError(f'Некорректное количество в {name}: {barcode}')
                chrt = by_barcode.get(barcode)
                if chrt is None:
                    raise StockOperationError(f'ШК {barcode} отсутствует или неоднозначен в каталоге {cab_key}; загрузка отменена')
                if chrt in result:
                    raise StockOperationError(f'Два ШК одного chrtID {chrt} в файле {name}; загрузка отменена')
                result[chrt] = qty
        finally:
            book.close()
        if not result:
            raise StockOperationError(f'Пустой складской файл {name}; загрузка отменена')
        targets[key] = result
        labels[key] = f'{wh.cabinet_name} — {wh.name}'
    return targets, labels


class StockManager:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.root = settings.base_dir / 'data' / 'stock_operations'
        self.previews = self.root / 'previews'
        self.records = self.root / 'records'
        self.lock = asyncio.Lock()
        self.active_previews: dict[int, str] = {}
        self.zero_authorized: set[tuple[int, str]] = set()

    def _record_path(self, ident: str) -> Path:
        if not ident.isalnum():
            raise StockOperationError('Некорректный ID операции')
        return self.records / f'{ident}.json'

    def _require_admin(self, actor: int) -> None:
        if actor not in self.settings.allowed_ids or actor not in self.settings.admin_ids:
            raise StockOperationError('Операция доступна только администратору')

    def _exclusions(self) -> dict[str, set[str]]:
        cabinets = self.settings.cabinets
        return read_exclusions(self.settings.base_dir / 'data' / 'exclusions.xlsx',
                               tuple(cab.key for cab in cabinets),
                               tuple(cab.seller_id for cab in cabinets))

    async def warehouse_choices(self, actor: int) -> dict[str, str]:
        self._require_admin(actor)
        async with aiohttp.ClientSession() as session:
            result = {}
            for cabinet in self.settings.cabinets:
                for warehouse in await WBClient(cabinet, session).get_warehouses():
                    result[_key(cabinet.key, warehouse.warehouse_id)] = f'{cabinet.name} — {warehouse.name}'
            return result

    def _preview_path(self, ident: str) -> Path:
        return self.previews / f'{ident}.json'

    def _cabinets(self, keys: set[str]):
        result = [cab for cab in self.settings.cabinets if cab.key in keys]
        if not result or {cab.key for cab in result} != keys:
            raise StockOperationError('Неизвестный кабинет')
        return result

    @staticmethod
    def _catalog(by_chrt: dict) -> tuple[dict[str, int | None], dict[int, str]]:
        by_sku: dict[str, int | None] = {}
        by_id: dict[int, str] = {}
        for chrt, variant in by_chrt.items():
            by_id[int(chrt)] = variant.skus[0]
            for sku in variant.skus:
                if sku in by_sku and by_sku[sku] != int(chrt):
                    by_sku[sku] = None  # fail closed for ambiguous SKU
                else:
                    by_sku[sku] = int(chrt)
        return by_sku, by_id

    async def preview(self, kind: str, actor: int, keys: set[str], files: list[Path] | None = None,
                      source_id: str | None = None, warehouse_keys: set[str] | None = None) -> dict:
        self._require_admin(actor)
        if self.lock.locked():
            raise StockOperationError("Идёт загрузка на WB. Дождитесь завершения")
        if kind not in {'distribution', 'zero', 'rollback'}:
            raise StockOperationError('Неизвестный тип операции')
        if kind == 'distribution' and not files:
            raise StockOperationError('Нет файлов распределения')
        if kind == 'rollback' and not source_id:
            raise StockOperationError('Нет резервной копии для отката')
        cabinets = self._cabinets(keys)
        targets: dict[str, dict[int, int]] = {}
        labels: dict[str, str] = {}
        skus: dict[str, dict[int, str]] = {}
        exclusions = self._exclusions()
        protected = {}
        async with aiohttp.ClientSession() as session:
            clients = {cab.key: WBClient(cab, session) for cab in cabinets}
            warehouses = {}
            catalogs = {}
            for cab in cabinets:
                client = clients[cab.key]
                whs = await client.get_warehouses()
                if not whs:
                    raise StockOperationError(f'Нет доступных FBS-складов: {cab.name}')
                for wh in whs:
                    warehouses[_key(cab.key, wh.warehouse_id)] = wh
                _barcodes, by_chrt = await client.get_catalog_variants()
                catalogs[cab.key], skus[cab.key] = self._catalog(by_chrt)
                protected[cab.key] = {int(chrt) for chrt, variant in by_chrt.items()
                                      if set(variant.skus) & exclusions.get(cab.key, set())}
                if not by_chrt:
                    raise StockOperationError(f'Не удалось получить каталог {cab.name}')
            if warehouse_keys is not None:
                if kind != 'zero' or not warehouse_keys or not warehouse_keys.issubset(warehouses):
                    raise StockOperationError('Некорректный выбор складов')
                warehouses = {key: value for key, value in warehouses.items() if key in warehouse_keys}
            if kind == 'distribution':
                targets, labels = load_distribution_targets(files or [], warehouses, catalogs)
            elif kind == 'zero':
                # The API does not list all warehouse chrtIDs without a catalog.
                # We query all current active catalog variants for each warehouse.
                for key, wh in warehouses.items():
                    cab_key = wh.cabinet_key
                    labels[key] = f'{wh.cabinet_name} — {wh.name}'
                    targets[key] = {chrt: 0 for chrt in skus[cab_key]}
            else:
                record = _load(self._record_path(source_id or ''))
                if record.get('kind') not in {'distribution', 'zero'} or record.get('rolled_back_at'):
                    raise StockOperationError('Эту операцию нельзя откатить')
                if record.get('status') not in {'verified', 'partial'}:
                    raise StockOperationError('Нет подтверждённых обновлений для отката')
                applied = record.get('applied', {})
                if not any(applied.values()):
                    raise StockOperationError('Нет успешно применённых позиций для отката')
                for key, rows in applied.items():
                    if key.split(':', 1)[0] not in keys or key in record.get('rolled_back_keys', []):
                        continue
                    if key not in warehouses:
                        raise StockOperationError(f'Склад {key} удалён; откат невозможен')
                    # Only restore positions that were actually verified as updated.
                    restored = set(record.get('restored', {}).get(key, []))
                    targets[key] = {int(chrt): int(record['before'][key][chrt]) for chrt in rows if chrt not in restored}
                    labels[key] = record['labels'][key]
                if not targets:
                    raise StockOperationError('Нет подтверждённых позиций для выбранных кабинетов')
            excluded_positions = []
            for key, rows in targets.items():
                for chrt in list(rows):
                    if chrt in protected[key.split(':', 1)[0]]:
                        excluded_positions.append(f'{key}:{chrt}')
                        del rows[chrt]
            before: dict[str, dict[int, int]] = {}
            if kind == 'zero':
                # Read every known active variant. Zero only the nonzero ones.
                for key, target in targets.items():
                    cab, warehouse = key.split(':')
                    old = await clients[cab].get_fbs_stocks(int(warehouse), target)
                    targets[key] = {chrt: 0 for chrt, qty in old.items() if qty > 0}
                    before[key] = {chrt: qty for chrt, qty in old.items() if qty > 0}
            else:
                for key, target in targets.items():
                    cab, warehouse = key.split(':')
                    before[key] = await clients[cab].get_fbs_stocks(int(warehouse), target)
            conflicts = []
            if kind == 'rollback':
                original = _load(self._record_path(source_id or ''))
                for key in list(targets):
                    for chrt in list(targets[key]):
                        if before[key][chrt] != int(original['after'][key][str(chrt)]):
                            conflicts.append(f'{key}:{chrt}')
                            targets[key].pop(chrt)
                            before[key].pop(chrt)
            data = {
                'id': uuid.uuid4().hex, 'kind': kind, 'actor': actor, 'created_at': utcnow(),
                'cabinets': sorted(keys), 'targets': {key: {str(c): q for c, q in rows.items()}
                                                     for key, rows in targets.items()},
                'before': {key: {str(c): q for c, q in rows.items()} for key, rows in before.items()},
                'labels': labels, 'skus': {cab: {str(c): sku for c, sku in mapping.items()}
                                          for cab, mapping in skus.items()},
                'source_id': source_id, 'conflicts': conflicts,
                'excluded_positions': excluded_positions,
                'exclusions': {key: sorted(value) for key, value in exclusions.items()},
                'summary': compare({key: {str(c): q for c, q in rows.items()} for key, rows in before.items()},
                                   {key: {str(c): q for c, q in rows.items()} for key, rows in targets.items()}),
            }
            _save(self._preview_path(data['id']), data)
            self.active_previews[actor] = data['id']
            return data

    def get_preview(self, ident: str, actor: int) -> dict:
        if not ident.isalnum():
            raise StockOperationError('Некорректный ID операции')
        path = self._preview_path(ident)
        if not path.is_file():
            raise StockOperationError('Предпросмотр не найден или устарел')
        data = _load(path)
        if data['actor'] != actor:
            raise StockOperationError('Предпросмотр принадлежит другому пользователю')
        if self.active_previews.get(actor) != ident:
            raise StockOperationError('Предпросмотр отменён или заменён новым')
        created = datetime.fromisoformat(data['created_at'])
        if (datetime.now(timezone.utc) - created).total_seconds() > 900:
            raise StockOperationError('Прошло более 15 минут. Повторите сравнение с WB')
        return data

    def list_records(self, limit: int | None = 10) -> list[dict]:
        records = []
        for path in self.records.glob('*.json'):
            try:
                records.append(_load(path))
            except (OSError, ValueError):
                continue
        return sorted(records, key=lambda r: r['created_at'], reverse=True)[:limit]

    def latest_rollback_source(self) -> dict | None:
        for item in self.list_records(None):
            if item.get('kind') in {'distribution', 'zero'} and any(item.get('applied', {}).values()) and not item.get('rolled_back_at'):
                return item
        return None

    def preview_report(self, data: dict) -> Path:
        target = self.previews / f"{data['id']}_comparison.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = 'Было и станет'
        sheet.append(['Кабинет / склад', 'ШК', 'chrtID', 'Было на WB', 'Станет', 'Разница'])
        for key, rows in sorted(data['targets'].items()):
            cab = key.split(':')[0]
            for chrt, new in sorted(rows.items(), key=lambda x: int(x[0])):
                old = data['before'][key][chrt]
                sheet.append([data['labels'][key], data['skus'].get(cab, {}).get(chrt, ''),
                              int(chrt), old, new, new - old])
        sheet.freeze_panes = 'A2'
        for col, width in {'A': 44, 'B': 22, 'C': 18, 'D': 16, 'E': 16, 'F': 16}.items():
            sheet.column_dimensions[col].width = width
        book.save(target)
        return target

    async def execute(self, ident: str, actor: int) -> dict:
        self._require_admin(actor)
        if self.lock.locked():
            raise StockOperationError('Уже выполняется другая операция с остатками')
        async with self.lock:
            preview = self.get_preview(ident, actor)
            if preview.get('exclusions') != {key: sorted(value) for key, value in self._exclusions().items()}:
                raise StockOperationError('Исключения изменились. Повторите сравнение с WB')
            if self._record_path(ident).exists():
                raise StockOperationError('Эта операция уже запускалась')
            if preview["kind"] == "zero" and (actor, ident) not in self.zero_authorized:
                raise StockOperationError("Аварийное обнуление не подтверждено словом ОБНУЛИТЬ")
            targets = {key: {int(chrt): int(q) for chrt, q in rows.items()}
                       for key, rows in preview['targets'].items()}
            before = {key: {int(chrt): int(q) for chrt, q in rows.items()}
                      for key, rows in preview['before'].items()}
            # Ensure latest operation is still the one chosen for rollback.
            if preview['kind'] == 'rollback':
                latest = self.latest_rollback_source()
                if latest is None or latest['id'] != preview['source_id']:
                    raise StockOperationError('Появилась новая загрузка; сделайте новый предпросмотр отката')
            record = dict(preview)
            record.update({'status': 'started', 'started_at': utcnow(), 'after': {},
                           'applied': {}, 'errors': {}, 'finished_at': None})
            # Record before ANY PUT so backups survive restarts and partial failure.
            _save(self._record_path(ident), record)
            async with aiohttp.ClientSession() as session:
                clients = {cab.key: WBClient(cab, session)
                           for cab in self._cabinets(set(preview['cabinets']))}
                # Abort entire operation if WB has changed since preview.
                try:
                    for key, expected in before.items():
                        cab, warehouse = key.split(':')
                        actual = await clients[cab].get_fbs_stocks(int(warehouse), expected)
                        if actual != expected:
                            raise StockOperationError(
                                f'Остатки склада {preview["labels"][key]} изменились после сравнения. '
                                'Ни один PUT не выполнен. Повторите предпросмотр.'
                            )
                except Exception as exc:
                    record['status'] = 'aborted'
                    record['errors']['preflight'] = str(exc)
                    record['finished_at'] = utcnow()
                    _save(self._record_path(ident), record)
                    raise
                for key, desired in targets.items():
                    cab, warehouse = key.split(':')
                    changed = {chrt: qty for chrt, qty in desired.items() if before[key][chrt] != qty}
                    record['after'][key] = {str(chrt): qty for chrt, qty in changed.items()}
                    record['applied'][key] = []
                    _save(self._record_path(ident), record)
                    for chunk in _chunks(changed):
                        sent_or_uncertain = False
                        try:
                            if preview['exclusions'] != {name: sorted(value) for name, value in self._exclusions().items()}:
                                raise StockOperationError('Исключения изменились во время загрузки')
                            # One more check before each PUT, to catch concurrent WB changes.
                            actual = await clients[cab].get_fbs_stocks(int(warehouse), chunk)
                            if any(actual[chrt] != before[key][chrt] for chrt in chunk):
                                raise StockOperationError('WB изменил остатки во время загрузки')
                            sent_or_uncertain = True
                            await clients[cab].put_fbs_stocks(int(warehouse), chunk)
                        except Exception as exc:
                            record['errors'].setdefault(key, []).append(str(exc))
                        # Verify even on timeout: PUT may have reached WB.
                        if not sent_or_uncertain:
                            _save(self._record_path(ident), record)
                            continue
                        try:
                            actual = await clients[cab].get_fbs_stocks(int(warehouse), chunk)
                            for chrt, qty in chunk.items():
                                if actual[chrt] == qty:
                                    record['applied'][key].append(str(chrt))
                                else:
                                    record['errors'].setdefault(key, []).append(
                                        f'chrtID {chrt}: ожидалось {qty}, WB вернул {actual[chrt]}')
                        except Exception as exc:
                            record['errors'].setdefault(key, []).append(f'Не удалось проверить: {exc}')
                        _save(self._record_path(ident), record)
            expected_count = sum(sum(1 for chrt, qty in rows.items() if before[key][chrt] != qty)
                                 for key, rows in targets.items())
            verified = sum(len(ids) for ids in record['applied'].values())
            record['status'] = ('verified' if verified == expected_count and not record['errors']
                                else 'partial')
            record['verified_count'] = verified
            record['expected_count'] = expected_count
            record['finished_at'] = utcnow()
            _save(self._record_path(ident), record)
            self.active_previews.pop(actor, None)
            self.zero_authorized.discard((actor, ident))
            if preview['kind'] == 'rollback':
                source = _load(self._record_path(preview['source_id']))
                # A rollback may select only one cabinet or skip conflicted chrtIDs.
                # Mark a warehouse restored only when ALL its applied IDs were restored.
                restored_keys = set(source.get('rolled_back_keys', []))
                restored = source.setdefault('restored', {})
                for key, ids in source.get('applied', {}).items():
                    restored[key] = sorted(set(restored.get(key, [])) | set(record['applied'].get(key, [])))
                    if ids and set(ids).issubset(restored[key]):
                        restored_keys.add(key)
                source['rolled_back_keys'] = sorted(restored_keys)
                if all(not ids or key in restored_keys for key, ids in source.get('applied', {}).items()):
                    source['rolled_back_at'] = utcnow()
                    source['rolled_back_by'] = actor
                _save(self._record_path(source['id']), source)
            return record
