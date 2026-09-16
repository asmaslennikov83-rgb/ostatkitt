from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from app.excel_io import write_warehouse_files
from app.models import DistributionLine, ProductVariant, Warehouse
from app.service import (
    _build_output_sku_maps,
    _build_sku_alias_groups,
    _expand_variants_for_aliases,
    _normalize_stock_by_alias_group,
)


def test_alternative_sku_of_same_chrtid_never_gets_zero_row():
    sku_old = "2040826951870"
    sku_input = "4600987014036"
    variant = ProductVariant("cab1", 777001, 123456, (sku_old, sku_input))
    actual = {"cab1": {sku_old: variant, sku_input: variant}}

    alias_root, groups = _build_sku_alias_groups(actual)
    normalized, excluded_roots, preferred, _merged = _normalize_stock_by_alias_group(
        {sku_input: 20}, set(), alias_root
    )
    assert normalized == {sku_input: 20}

    expanded = _expand_variants_for_aliases(actual, groups)
    assert expanded["cab1"][sku_old].chrt_id == expanded["cab1"][sku_input].chrt_id

    all_barcodes, canonical = _build_output_sku_maps(
        actual_by_cabinet=actual,
        groups=groups,
        alias_root=alias_root,
        excluded_roots=excluded_roots,
        preferred_input_by_root=preferred,
        kit_barcodes=set(),
    )
    assert all_barcodes["cab1"] == {sku_input}
    assert canonical["cab1"][sku_old] == sku_input
    assert canonical["cab1"][sku_input] == sku_input

    with TemporaryDirectory() as td:
        td = Path(td)
        template = td / "template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["Баркод", "Количество"])
        wb.save(template)

        wh = Warehouse("cab1", "Кабинет", 10, "Склад")
        line = DistributionLine(barcode=sku_input, source_qty=20, allocations={("cab1", 10): 20})
        files = write_warehouse_files(
            template, td / "out", [wh], [line], all_barcodes,
            output_barcode_by_cabinet=canonical,
        )
        out_wb = load_workbook(files[0], data_only=True)
        rows = list(out_wb.active.iter_rows(min_row=2, values_only=True))
        assert rows == [(sku_input, 20)]
        assert (sku_old, 0) not in rows


def test_multiple_input_aliases_are_aggregated_by_variant():
    a, b = "111", "222"
    variant = ProductVariant("cab1", 99, 1, (a, b))
    actual = {"cab1": {a: variant, b: variant}}
    alias_root, _groups = _build_sku_alias_groups(actual)
    normalized, _excluded, _preferred, merged = _normalize_stock_by_alias_group(
        {a: 4, b: 6}, set(), alias_root
    )
    assert sum(normalized.values()) == 10
    assert len(normalized) == 1
    assert any(set(v) == {a, b} for v in merged.values())
