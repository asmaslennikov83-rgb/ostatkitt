from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from app.excel_io import write_warehouse_files
from app.models import DistributionLine, Warehouse


def test_full_catalog_zero_export():
    with TemporaryDirectory() as td:
        root = Path(td)
        template = root / "template.xlsx"
        out = root / "out"
        wb = Workbook()
        ws = wb.active
        ws.append(["Баркод", "Количество"])
        wb.save(template)

        warehouses = [Warehouse("cab1", "Кабинет 1", 101, "СПб")]
        line = DistributionLine(barcode="111", source_qty=5)
        line.allocations[("cab1", 101)] = 3

        files = write_warehouse_files(
            template,
            out,
            warehouses,
            [line],
            {"cab1": {"111", "222", "333"}},
        )
        result = load_workbook(files[0], data_only=True).active
        rows = {str(result.cell(r, 1).value): int(result.cell(r, 2).value) for r in range(2, result.max_row + 1)}
        assert rows == {"111": 3, "222": 0, "333": 0}


if __name__ == "__main__":
    test_full_catalog_zero_export()
    print("OK")
