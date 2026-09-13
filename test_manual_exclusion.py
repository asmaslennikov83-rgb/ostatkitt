from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import Workbook, load_workbook

from app.excel_io import read_input_xlsx, write_warehouse_files
from app.models import DistributionLine, Warehouse

with TemporaryDirectory() as td:
    td = Path(td)
    inp = td / 'input.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.append(['Баркод', 'Количество'])
    ws.append(['111', 5])
    ws.append(['222', '!'])
    ws.append(['333', 0])
    wb.save(inp)

    stock, excluded = read_input_xlsx(inp)
    assert stock == {'111': 5, '333': 0}, stock
    assert excluded == {'222'}, excluded

    template = td / 'template.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.append(['Баркод', 'Количество'])
    wb.save(template)

    wh = Warehouse('cab1', 'Кабинет', 1, 'Склад')
    lines = [DistributionLine(barcode='111', source_qty=5, allocations={('cab1', 1): 5})]
    # excluded 222 intentionally removed before export; 333 remains and should be zeroed
    files = write_warehouse_files(template, td/'out', [wh], lines, {'cab1': {'111','333'}})
    out = load_workbook(files[0], read_only=True).active
    rows = {str(r[0].value): r[1].value for r in out.iter_rows(min_row=2) if r[0].value is not None}
    assert rows == {'111': 5, '333': 0}, rows
    assert '222' not in rows
print('OK manual exclusion')
