from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from app.exclusions import read_exclusions
from app.excel_io import write_warehouse_files
from app.models import DistributionLine, ProductVariant, Warehouse
from app.service import (
    _apply_cabinet_barcode_exclusions,
    _build_output_sku_maps,
    _build_sku_alias_groups,
    _expand_variants_for_aliases,
    _normalize_stock_by_alias_group,
)


def test_exclusions_template_is_per_cabinet():
    with TemporaryDirectory() as td:
        path = Path(td) / "exclusions.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["ШК", "ID кабинета 1", "ID кабинета 2"])
        ws.append(["4600987012193", "", "seller-2"])
        ws.append(["111", "seller-1", ""])
        wb.save(path)
        data = read_exclusions(path, ("cab1", "cab2"), ("seller-1", "seller-2"))
        assert data == {"cab1": {"111"}, "cab2": {"4600987012193"}}


def test_forbidden_alias_never_written_but_allowed_alias_keeps_group():
    forbidden = "4600987012193"
    allowed = "4600987012223"
    v1 = ProductVariant("cab1", 700, 1, (forbidden, allowed))
    v2 = ProductVariant("cab2", 800, 2, (forbidden, allowed))
    raw = {
        "cab1": {forbidden: v1, allowed: v1},
        "cab2": {forbidden: v2, allowed: v2},
    }
    alias_root, groups = _build_sku_alias_groups(raw)
    effective = _apply_cabinet_barcode_exclusions(raw, {"cab1": set(), "cab2": {forbidden}})
    expanded = _expand_variants_for_aliases(effective, groups)

    # Physical stock can still participate in cab2 because an allowed alias exists.
    assert forbidden in expanded["cab2"]
    assert expanded["cab2"][forbidden].skus == (allowed,)

    normalized, excluded_roots, preferred, _ = _normalize_stock_by_alias_group(
        {forbidden: 25}, set(), alias_root
    )
    assert normalized == {forbidden: 25}

    all_barcodes, canonical = _build_output_sku_maps(
        actual_by_cabinet=effective,
        groups=groups,
        alias_root=alias_root,
        excluded_roots=excluded_roots,
        preferred_input_by_root=preferred,
        kit_barcodes=set(),
    )
    assert forbidden not in all_barcodes["cab2"]
    assert all_barcodes["cab2"] == {allowed}
    assert canonical["cab2"][forbidden] == allowed

    with TemporaryDirectory() as td:
        td = Path(td)
        template = td / "template.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["Баркод", "Количество"])
        wb.save(template)
        wh = Warehouse("cab2", "Кабинет 2", 10, "Центральный регион")
        line = DistributionLine(barcode=forbidden, source_qty=25, allocations={("cab2", 10): 3})
        files = write_warehouse_files(
            template, td / "out", [wh], [line], all_barcodes,
            output_barcode_by_cabinet=canonical,
        )
        rows = list(load_workbook(files[0], data_only=True).active.iter_rows(min_row=2, values_only=True))
        assert rows == [(allowed, 3)]
        assert (forbidden, 3) not in rows
        assert (forbidden, 0) not in rows


def test_if_all_aliases_forbidden_group_not_distributed_to_cabinet():
    a, b = "111", "222"
    v = ProductVariant("cab2", 800, 2, (a, b))
    raw = {"cab2": {a: v, b: v}}
    alias_root, groups = _build_sku_alias_groups(raw)
    effective = _apply_cabinet_barcode_exclusions(raw, {"cab2": {a, b}})
    expanded = _expand_variants_for_aliases(effective, groups)
    assert expanded["cab2"] == {}

    normalized, excluded_roots, preferred, _ = _normalize_stock_by_alias_group({a: 10}, set(), alias_root)
    all_barcodes, canonical = _build_output_sku_maps(
        effective, groups, alias_root, excluded_roots, preferred, set()
    )
    assert all_barcodes["cab2"] == set()
    assert canonical["cab2"] == {}


def test_wrong_seller_id_is_rejected():
    with TemporaryDirectory() as td:
        path = Path(td) / "exclusions.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["ШК", "ID кабинета 1", "ID кабинета 2"])
        ws.append(["111", "WRONG", ""])
        wb.save(path)
        try:
            read_exclusions(path, ("cab1", "cab2"), ("seller-1", "seller-2"))
        except ValueError as exc:
            assert "ожидался ID продавца seller-1" in str(exc)
        else:
            raise AssertionError("wrong seller ID must be rejected")
